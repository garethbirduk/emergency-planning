"""Regression suite from build step 2.4: roles & personalised checklists
(AC-R1–R3).

Sample: Alex Sample = executor + admin, organiser; Chris Sample = helper +
local, medical; `comms` held by nobody (fallback path)."""

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SCRATCH = Path(tempfile.mkdtemp(prefix="crisisplan-verify24-"))
sys.path.insert(0, str(REPO / "tools"))

failures = []


def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(name)


def run_cli(src, out_dir):
    r = subprocess.run(
        [sys.executable, str(REPO / "tools" / "generate.py"), "-f", str(src),
         "--date", "2030-01-15", "--out", str(out_dir)],
        capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr)
        raise SystemExit("CLI failed")


out = SCRATCH / "run-24"
run_cli(REPO / "samples" / "pat.sample.xlsx", out)
md = (out / "checklist.md").read_text(encoding="utf-8")

print("AC-R1: every task assigned; unheld role visible")
task_lines = [ln for ln in md.splitlines() if ln.startswith("- [ ] ")]
unowned = [ln for ln in task_lines
           if not re.match(r"- \[.\] (?:📜 )?\*\*[A-Z]+ \(", ln)]
check("every master task line carries an assignee",
      task_lines and not unowned,
      f"{len(task_lines)} tasks" + (f"; unowned: {unowned[:2]}" if unowned else ""))
check("comms fallback note present and names the executor",
      "role 'comms': held by nobody on the household tab — its tasks "
      "default to Alex Sample" in md)
check("no other fallback notes (all other roles held)",
      md.count("held by nobody") == 1)

print("Role routing spot-checks")


def owner_of(title_fragment):
    for ln in task_lines:
        if title_fragment in ln:
            m = re.match(r"- \[.\] (?:📜 )?\*\*[A-Z]+ \((.+?)\):\*\*", ln)
            return m.group(1) if m else None
    return None


check("medical -> Chris (medical certificate)",
      owner_of("medical certificate") == "Chris Sample",
      str(owner_of("medical certificate")))
check("local -> Chris (tell close family)",
      owner_of("Tell close family") == "Chris Sample")
check("organiser -> Alex (funeral director)",
      owner_of("funeral director") == "Alex Sample")
check("admin -> Alex (register the death)",
      owner_of("Register the death") == "Alex Sample")
check("comms fallback -> Alex (invites)",
      owner_of("Invite The Sample cousins") == "Alex Sample")
check("organ donation wish -> Chris (medical, day 0)",
      owner_of("organ_donation") == "Chris Sample")

print("AC-R2: per-person checklists")
alex_p = out / "checklist-alex-sample.md"
chris_p = out / "checklist-chris-sample.md"
check("both personal checklists written",
      alex_p.exists() and chris_p.exists(),
      ", ".join(p.name for p in out.glob("checklist-*.md")))
if alex_p.exists() and chris_p.exists():
    alex = alex_p.read_text(encoding="utf-8")
    chris = chris_p.read_text(encoding="utf-8")
    n_alex = alex.count("- [ ] ")
    n_chris = chris.count("- [ ] ")
    check("union of personal checklists equals the master",
          n_alex + n_chris == len(task_lines),
          f"{n_alex} + {n_chris} vs {len(task_lines)}")
    check("personal files are for the right person",
          "— Alex Sample" in alex.splitlines()[0]
          and "Invite The Sample cousins" in alex
          and "Tell close family" in chris)

print("AC-R3: role edits change assignees, never UIDs")
from openpyxl import load_workbook  # noqa: E402


def uids(path):
    text = path.read_bytes().decode("utf-8").replace("\r\n ", "")
    return set(re.findall(r"^UID:(.+)$", text, re.M))


copy = SCRATCH / "pat.roles.xlsx"
shutil.copyfile(REPO / "samples" / "pat.sample.xlsx", copy)
wb = load_workbook(copy)
ws = wb["household"]
assert ws.cell(row=4, column=1).value == "Chris Sample"
ws.cell(row=4, column=4).value = "local"  # drop 'medical' from Chris
wb.save(copy)
out2 = SCRATCH / "run-24b"
run_cli(copy, out2)
md2 = (out2 / "checklist.md").read_text(encoding="utf-8")
check("UID set identical after the role edit",
      uids(out / "plan.ics") == uids(out2 / "plan.ics"))
check("medical tasks reassigned by fallback, with a visible note",
      "role 'medical': held by nobody" in md2
      and "- [ ] **MEDICAL (Alex Sample):** Obtain the medical "
          "certificate of cause of death" in md2)

print()
if failures:
    raise SystemExit(f"FAILED: {failures}")
print(f"All step-2.4 checks passed ({len(task_lines)} tasks assigned).")
