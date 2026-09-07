"""Regression suite from build step 4.1: other households.

One engine, data-driven: a blank template generated outside the repo,
filled for a sole-name-house household (all dummy names), produces
the correct plan — including the sole-owner house asymmetry the design
was born from (AC-C3 on a second household's file).
"""

import json
import os
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SCRATCH = Path(tempfile.mkdtemp(prefix="crisisplan-verify41-"))
sys.path.insert(0, str(REPO / "tools"))

failures = []


def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(name)


def run(script, *args, env=None):
    return subprocess.run(
        [sys.executable, str(REPO / "tools" / script), *map(str, args)],
        capture_output=True, text=True, encoding="utf-8",
        errors="replace", env=env)


print("Blank template generation")
blank = SCRATCH / "lee.xlsx"
r = run("make_template.py", "--blank", "--out", blank)
check("blank template written outside the repo",
      r.returncode == 0 and blank.exists(), r.stderr.strip()[:200])
from openpyxl import load_workbook  # noqa: E402
from reader import read_workbook  # noqa: E402
from tasks import derive_tasks  # noqa: E402

wb = load_workbook(blank, read_only=True)
sheets = wb.sheetnames
wb.close()
sample_sheets = load_workbook(REPO / "samples" / "pat.sample.xlsx",
                              read_only=True).sheetnames
check("same tabs and order as the sample template",
      sheets == sample_sheets)
model = read_workbook(blank)
check("every module tab present but empty (no dummy rows)",
      not model.modules and len(model.empty_tabs) == 18,
      f"{len(model.empty_tabs)} empty tabs")
r = run("make_template.py", "--blank", "--out",
        REPO / "tools" / "x.xlsx")
check("repo-tree --out refused",
      r.returncode != 0 and "REFUSED" in (r.stderr + r.stdout))

print("Blank workbook still runs end-to-end")
out_dir = SCRATCH / "run-blank"
r = run("generate.py", "-f", blank, "--date", "2030-06-01",
        "--out", out_dir)
check("run OK on an empty household",
      r.returncode == 0 and (out_dir / "checklist.md").exists())
md = (out_dir / "checklist.md").read_text(encoding="utf-8")
check("base tasks only, nothing invented",
      md.count("- [ ] ") == 7 and "`base/" in md)

print("sole-name-house household (dummy data, same engine)")
lee = SCRATCH / "lee-filled.xlsx"
wb = load_workbook(blank)


def fill(tab, rows):
    ws = wb[tab]
    header = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}
    for r_off, row in enumerate(rows):
        for col_name, value in row.items():
            ws.cell(row=2 + r_off, column=header[col_name]).value = value


fill("household", [
    {"name": "Lee Example", "role": "planholder",
     "relationship": "self"},
    {"name": "Vic Example", "relationship": "spouse",
     "crisis_roles": "organiser, medical, local, comms"},
    {"name": "Nephew Example", "role": "executor",
     "crisis_roles": "admin"},
])
fill("property", [
    # The asymmetry the whole design started from: the house in the
    # planholder's SOLE name, spouse living in it.
    {"property_label": "The bungalow", "joint?": "sole",
     "owner": "Lee", "mortgage?": "N", "insurer": "Example Insurance"},
])
fill("building_societies", [
    {"society": "Example Mutual BS", "account_label": "Passbook saver",
     "joint?": "sole", "owner": "Lee",
     "location_of_details": "Passbook drawer"},
])
wb.save(lee)

tasks, notes = derive_tasks(read_workbook(lee), date(2030, 6, 1))
keys = {t.plan_key for t in tasks}
check("sole house -> probate path (secure + transfer, no DJP)",
      any(k.endswith("the-bungalow/secure-property") for k in keys)
      and any(k.endswith("the-bungalow/transfer-or-sell") for k in keys)
      and not any("land-registry-djp" in k for k in keys))
check("sole BS account -> probate path (date-of-death balance)",
      any(k.endswith("/dod-balance") for k in keys))
by_key = {t.plan_key: t for t in tasks}
check("executor nephew gets admin, spouse gets organiser/medical",
      by_key["base/register-death"].assignee == "Nephew Example"
      and by_key["base/funeral-director"].assignee == "Vic Example"
      and by_key["base/verify-death"].assignee == "Vic Example")

print("Unheld roles fall back to the spouse, found by relationship")
# 'role' is legal standing only now — Vic Example is the fallback because their
# relationship says spouse, not because of any role value.
wb = load_workbook(lee)
ws = wb["household"]
header = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}
ws.cell(row=4, column=header["role"]).value = ""          # no executor now
ws.cell(row=4, column=header["crisis_roles"]).value = ""  # admin unheld
nofb = SCRATCH / "lee-no-executor.xlsx"
wb.save(nofb)
tasks3, notes3 = derive_tasks(read_workbook(nofb), date(2030, 6, 1))
by_key3 = {t.plan_key: t for t in tasks3}
check("admin defaults to the spouse by relationship, with a note",
      by_key3["base/register-death"].assignee == "Vic Example"
      and any("default to Vic Example" in n for n in notes3))

print("Flip to joint -> survivor path (AC-C3 on this household)")
wb = load_workbook(lee)
ws = wb["property"]
header = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}
ws.cell(row=2, column=header["joint?"]).value = "joint"
ws.cell(row=2, column=header["tenancy"]).value = "joint_tenants"
joint = SCRATCH / "lee-joint.xlsx"
wb.save(joint)
tasks2, _ = derive_tasks(read_workbook(joint), date(2030, 6, 1))
keys2 = {t.plan_key for t in tasks2}
check("joint house -> DJP survivor path, probate tasks gone",
      any("land-registry-djp" in k for k in keys2)
      and not any(k.endswith("the-bungalow/secure-property")
                  for k in keys2))

print("Default output lands beside a real workbook")
r = run("generate.py", "-f", lee, "--date", "2030-06-01")
plan_dir = SCRATCH   # the workbook's own folder
check("no --out -> plan files land alongside the workbook",
      r.returncode == 0 and (plan_dir / "checklist.md").exists()
      and (plan_dir / "google-preview.md").exists(),
      str(plan_dir))
ghost = plan_dir / "checklist-ghost-person.md"
ghost.write_text("stale", encoding="utf-8")
run("generate.py", "-f", lee, "--date", "2030-06-01")
check("stale personal checklists cleaned on regeneration",
      not ghost.exists())

print("Per-household google config (--google-config)")
cfg_path = SCRATCH / "google-lee.json"
state = SCRATCH / "gcal-lee.json"
env = dict(os.environ, CRISISPLAN_GOOGLE_CONFIG=str(cfg_path),
           CRISISPLAN_GCAL_FAKE=str(state))
run("gsetup.py", "--emails", "", "--household", "Lee Example", env=env)
env.pop("CRISISPLAN_GOOGLE_CONFIG")
r = run("generate.py", "-f", lee, "--date", "2030-06-01",
        "--out", SCRATCH / "run-lee", "--mode", "test",
        "--google-config", cfg_path, env=env)
cfg = json.loads(cfg_path.read_text(encoding="utf-8"))
events = json.loads(state.read_text(encoding="utf-8"))["events"].get(
    cfg["test_calendar_id"], {})
check("--google-config routes this household to its own calendars",
      r.returncode == 0 and "loaded google-lee.json" in r.stdout
      and len(events) > 0, f"{len(events)} events")

print()
if failures:
    raise SystemExit(f"FAILED: {failures}")
print(f"All step-4.1 checks passed ({len(tasks)} tasks for the second "
      "household).")
