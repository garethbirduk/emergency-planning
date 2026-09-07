"""Regression suite from build step 2.1: every module emits sensible tasks
from sample rows, and the secret guard rejects (without echoing) rows that
look like they contain real secrets. Sample data only."""

import shutil
import sys
import tempfile
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "tools"))
from reader import read_workbook  # noqa: E402
from tasks import derive_tasks    # noqa: E402

SAMPLE = REPO / "samples" / "pat.sample.xlsx"
SCRATCH = Path(tempfile.mkdtemp(prefix="crisisplan-verify21-"))
TRIGGER = date(2030, 1, 15)

failures = []


def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(name)


# --- every wired module emits from the sample --------------------------------
print("Module coverage from the sample workbook")
model = read_workbook(SAMPLE)
tasks, notes = derive_tasks(model, TRIGGER)
by_module = {}
for t in tasks:
    by_module.setdefault(t.module, []).append(t)

EXPECT_TASKS = ["base", "banks", "building_societies", "cash_isas",
                "share_portfolios", "pensions", "property", "other_assets",
                "digital", "funeral_contacts", "wishes", "key_documents",
                "access_pointers"]
for mod in EXPECT_TASKS:
    check(f"{mod} emits tasks", len(by_module.get(mod, [])) > 0,
          f"{len(by_module.get(mod, []))} task(s)")
# shares_isas: the only sample row is Alex's -> skip note, not tasks (AC-C3).
check("shares_isas row becomes a skip note (owner is not the deceased)",
      any(n.startswith("shares_isas/") for n in notes))
check("sample 'POINTER ONLY — never the password itself' does NOT trip "
      "the secret guard",
      not any("contains an actual secret" in n for n in notes))

# spot-checks on content
digital = {t.plan_key: t for t in by_module.get("digital", [])}
check("digital: transfer action on the domain",
      any(k.startswith("digital/example-com-domain/") and
          t.title.startswith("Transfer") for k, t in digital.items()))
wishes = by_module.get("wishes", [])
organ = [t for t in wishes if "organ_donation" in t.title]
check("wishes: organ donation lands on day 0 (time-critical)",
      organ and organ[0].due == TRIGGER)
lpa = [t for t in by_module.get("key_documents", [])
       if "LPA" in t.title]
check("key_documents: LPA task carries the ends-at-death warning",
      lpa and any("ends at death" in d for d in lpa[0].details))
cert_count = sum(1 for t in tasks if t.needs_cert)
check("death-certificate steps flagged across the plan",
      cert_count >= 8, f"{cert_count} flagged")
check("bank notify needs a cert; funeral contact invite does not",
      next(t for t in by_module["banks"]
           if t.plan_key.endswith("/notify")).needs_cert
      and not any(t.needs_cert for t in by_module["funeral_contacts"]))
register = next(t for t in tasks
                if t.plan_key == "base/register-death")
check("copies-to-order estimate computed from the flagged count",
      any(str(cert_count) + " step(s)" in d for d in register.details)
      and any("certified copies" in d for d in register.details))

# --- secret guard on a doctored copy -----------------------------------------
print("Secret guard (doctored scratch copy)")
from openpyxl import load_workbook  # noqa: E402

copy = SCRATCH / "pat.secret-test.xlsx"
shutil.copyfile(SAMPLE, copy)
wb = load_workbook(copy)
ws = wb["access_pointers"]
ws.cell(row=5, column=1, value="Safe combination")
ws.cell(row=5, column=2, value="PIN is 4321")
wb.save(copy)

tasks2, notes2 = derive_tasks(read_workbook(copy), TRIGGER)
warn = [n for n in notes2 if "contains an actual secret" in n]
check("secret-looking row rejected with a warning", len(warn) == 1,
      warn[0] if warn else "no warning emitted")
check("the secret itself is NOT echoed in any output",
      all("4321" not in n for n in notes2)
      and all("4321" not in (t.title + " ".join(t.details)) for t in tasks2))
check("rejected row emits no tasks",
      not any("safe-combination" in t.plan_key for t in tasks2))
check("other access_pointers rows unaffected",
      sum(t.module == "access_pointers" for t in tasks2)
      == sum(t.module == "access_pointers" for t in tasks))

print()
if failures:
    raise SystemExit(f"FAILED: {failures}")
print(f"All step-2.1 checks passed ({len(tasks)} tasks, {len(notes)} notes "
      "from the sample).")
