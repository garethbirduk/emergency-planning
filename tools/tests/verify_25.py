"""Regression suite from build step 2.5: funeral, professionals,
life_insurance modules + chronological tab order (AC-C1 extension).

Sample rows (set in make_template.py, all default generic visibility):
  funeral #1         Example Funerals Ltd — prepaid cremation package
  professionals #1   Example & Co Solicitors (solicitor)
  professionals #2   Sample & Partners Accountancy (accountant)
  life_insurance #1  Example Life Assurance plc — term policy (no trust)
  life_insurance #2  Sample Employer Ltd — death in service (nominated)
"""

import subprocess
import sys
import shutil
import tempfile
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SAMPLE = REPO / "samples" / "pat.sample.xlsx"
SCRATCH = Path(tempfile.mkdtemp(prefix="crisisplan-verify25-"))
sys.path.insert(0, str(REPO / "tools"))

failures = []


def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(name)


def run_cli(src, out_dir):
    r = subprocess.run(
        [sys.executable, str(REPO / "tools" / "generate.py"), "-f", str(src),
         "--date", "2030-01-15", "--out", str(out_dir)],
        capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr)
        raise SystemExit("CLI failed")


def parse_ics(path):
    text = Path(path).read_bytes().decode("utf-8").replace("\r\n ", "")
    events, cur = {}, None
    for line in text.split("\r\n"):
        if line == "BEGIN:VEVENT":
            cur = {}
        elif line == "END:VEVENT":
            events[cur["UID"]] = cur
            cur = None
        elif cur is not None and ":" in line:
            prop, value = line.split(":", 1)
            cur[prop.split(";")[0]] = value
    return events


out = SCRATCH / "run-25"
run_cli(SAMPLE, out)
md = (out / "checklist.md").read_text(encoding="utf-8")
ics = (out / "plan.ics").read_bytes().decode("utf-8").replace("\r\n ", "")
events = parse_ics(out / "plan.ics")

print("Funeral override: pre-arranged provider replaces the base task")
check("base/funeral-director absent everywhere",
      "base/funeral-director" not in md and "base/funeral-director"
      not in ics)
check("override note present (never silent)",
      "the generic 'choose and contact a funeral director' task is "
      "replaced" in md)
check("contact-provider task on day 2, organiser -> Alex",
      "- [ ] **ORGANISER (Alex Sample):** Contact the pre-arranged "
      "funeral director: Example Funerals Ltd (funeral arrangement #1) "
      "→ `funeral/example-funerals-ltd/contact-provider`" in md
      and events.get("funeral/1/contact-provider", {}).get("DTSTART")
      == "20300117")
check("prepaid + package + venue details present",
      "already paid for" in md and "Package: Simple cremation package"
      in md and "Venue: Sampletown Crematorium" in md)

print("Professionals: type-driven tasks")
check("solicitor: notify day 3 + engage-probate day 30",
      events.get("professionals/1/notify", {}).get("DTSTART") == "20300118"
      and events.get("professionals/1/engage-probate",
                     {}).get("DTSTART") == "20300214")
check("solicitor notify points at key_documents",
      "They may hold the will or deeds" in md)
check("accountant: notify day 7 + final-tax day 60",
      events.get("professionals/2/notify", {}).get("DTSTART") == "20300122"
      and events.get("professionals/2/final-tax", {}).get("DTSTART")
      == "20300316")
check("accountant task covers tax to date of death",
      "settle tax affairs to the date of death" in md)

print("Life insurance: policies and death-in-service")
check("term policy: start-claim day 7 + chase day 45",
      events.get("life_insurance/1/start-claim", {}).get("DTSTART")
      == "20300122"
      and events.get("life_insurance/1/chase-claim", {}).get("DTSTART")
      == "20300301")
check("no-trust policy warns it may need probate",
      "Check whether the policy is written in trust" in md)
check("death-in-service: employer notify day 3 + claim day 30",
      events.get("life_insurance/2/notify", {}).get("DTSTART")
      == "20300118"
      and events.get("life_insurance/2/claim-benefits", {}).get("DTSTART")
      == "20300214")
check("death-in-service carries the ask-HR-anyway nudge",
      "classic benefit nobody knew about" in md)
check("expression of wish (Y) -> pays outside the estate",
      "benefit should pay to the nominee" in md)

print("Redaction: new tabs behave like the rest (default generic)")
new_names = ["Example Funerals", "Sampletown Crematorium", "Example & Co",
             "Sample & Partners", "Example Life Assurance",
             "Sample Employer"]
check("no new real name anywhere in the .ics",
      not any(n in ics for n in new_names),
      ", ".join(n for n in new_names if n in ics))
check("generic external titles used",
      events.get("funeral/1/contact-provider", {}).get("SUMMARY")
      == "Contact the pre-arranged funeral director: funeral "
         "arrangement #1"
      and events.get("professionals/1/engage-probate", {}).get("SUMMARY")
      == "Engage professional contact #1 for probate\\, or decide to "
         "handle it yourselves")   # \, — raw .ics TEXT escaping
check("legend decodes the new rows",
      "`funeral #1` = Example Funerals Ltd" in md
      and "`life_insurance #2` = Sample Employer Ltd" in md)

print("Override flips off: empty funeral tab restores the base task")
from openpyxl import load_workbook  # noqa: E402
from reader import read_workbook  # noqa: E402
from tasks import derive_tasks  # noqa: E402

copy = SCRATCH / "pat.no-funeral.xlsx"
shutil.copyfile(SAMPLE, copy)
wb = load_workbook(copy)
ws = wb["funeral"]
header = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}
for col_name, col in header.items():
    if col_name not in ("id", "online_visibility", "ai_visible?"):
        ws.cell(row=2, column=col).value = None
wb.save(copy)

tasks, notes = derive_tasks(read_workbook(copy), date(2030, 1, 15))
keys = {t.plan_key for t in tasks}
check("base/funeral-director back when the funeral tab is empty",
      "base/funeral-director" in keys)
check("no funeral-tab tasks and no override note",
      not any(k.startswith("funeral/") for k in keys)
      and not any("is replaced" in n for n in notes))

print("Builder branches not in the sample (synthetic rows)")
from openpyxl import Workbook  # noqa: E402

synth = SCRATCH / "synth.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "professionals"
ws.append(["name", "type", "contact_hint", "notes"])
ws.append(["Example Wealth Advice", "financial_adviser", "", ""])
ws.append(["Example Kennels", "other", "", ""])
ws2 = wb.create_sheet("life_insurance")
ws2.append(["provider", "policy_label", "type", "owner",
            "nomination_in_place?", "location_of_details", "notes"])
ws2.append(["Example Life Assurance plc", "Alex's own policy", "term",
            "Alex", "N", "", ""])
ws2.append(["Dup Insurance", "Dup Insurance", "term", "Pat", "N", "", ""])
ws2.append(["Zurich", "", "term", "Pat", "N", "", ""])   # label optional
ws3 = wb.create_sheet("household")
ws3.append(["name", "relationship", "role", "crisis_roles", "invited?",
            "region", "notes"])
ws3.append(["Pat Sample", "self", "planholder", "", "Y", "", ""])
wb.save(synth)

tasks, notes = derive_tasks(read_workbook(synth), date(2030, 1, 15))
keys = {t.plan_key for t in tasks}
check("financial_adviser: notify + asset-list",
      any(k.endswith("example-wealth-advice/notify") for k in keys)
      and any(k.endswith("example-wealth-advice/asset-list")
              for k in keys))
check("unknown type: plain notify only",
      any(k.endswith("example-kennels/notify") for k in keys)
      and sum(1 for k in keys if "example-kennels" in k) == 1)
check("life policy owned by someone else -> skip note, no tasks",
      not any("alex-s-own-policy" in k for k in keys)
      and any("sole asset owned by Alex" in n for n in notes))
check("duplicate key parts collapse (no zurich-zurich keys)",
      "life_insurance/dup-insurance/start-claim" in keys)
check("blank policy_label: key and title are the provider alone",
      "life_insurance/zurich/start-claim" in keys
      and any(t.plan_key == "life_insurance/zurich/start-claim"
              and t.title.startswith(
                  "Notify Zurich (life-insurance policy #")
              and " — " not in t.title for t in tasks))

print("Chronological tab order (presentational)")
wb = load_workbook(SAMPLE, read_only=True)
order = wb.sheetnames
wb.close()
check("sheets run people -> wishes -> funeral -> docs -> ... -> digital",
      order == ["instructions", "household", "preparation",
                "wishes", "funeral",
                "key_documents", "access_pointers", "funeral_contacts",
                "professionals", "banks", "building_societies",
                "cash_isas", "shares_isas", "share_portfolios",
                "pensions", "property", "other_assets", "life_insurance",
                "digital"],
      " -> ".join(order))

print()
if failures:
    raise SystemExit(f"FAILED: {failures}")
print(f"All step-2.5 checks passed ({len(events)} calendar events).")
