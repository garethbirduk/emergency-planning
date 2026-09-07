"""Regression suite from build step 5.3: the preparation plan
("what to do before someone dies" — SPEC §16 first slice).

Declared rows from the preparation tab + gaps the tool notices; a
local-only markdown, never a task, never on any calendar.
"""

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SAMPLE = REPO / "samples" / "pat.sample.xlsx"
SCRATCH = Path(tempfile.mkdtemp(prefix="crisisplan-verify53-"))
sys.path.insert(0, str(REPO / "tools"))

from prep import derive_gaps, render_preparation  # noqa: E402
from reader import read_workbook  # noqa: E402

failures = []


def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(name)


print("preparation.md written by a normal run")
out = SCRATCH / "run-53"
r = subprocess.run(
    [sys.executable, str(REPO / "tools" / "generate.py"), "-f",
     str(SAMPLE), "--date", "2030-01-15", "--out", str(out)],
    capture_output=True, text=True, encoding="utf-8", errors="replace")
check("run OK with a preparation health line",
      r.returncode == 0 and "preparation" in r.stdout
      and "completeness gap(s)" in r.stdout)
md = (out / "preparation.md").read_text(encoding="utf-8")

print("Declared rows render by kind and status")
check("actions carry priority, owner and status",
      "- [ ] **NOW** Make/update the wills — *Pat* *(in progress)*" in md
      and "- [ ] **SOON** Interview candidates — *Alex*" in md)
check("groups render as subsections, done sub-steps ticked",
      "### Arrange cleaner" in md
      and "- [x] **SOON** Advertise for a cleaner — *Alex*" in md)
check("closed items out of the plan, named in the footer",
      "Consider a stairlift" not in md.split("*Marked not needed")[0]
      and "*Marked not needed: Home accessibility: Consider a "
          "stairlift.*" in md)
check("preferences listed, not checkboxed",
      "- No stairlift — wants to stay mobile — Discussed and agreed "
      "2026" in md)
check("status lives in the workbook, not in md ticks",
      "Track status in the spreadsheet" in md)

print("Completeness gaps derived from the rest of the workbook")
check("missing pension nomination noticed",
      "Lifetime annuity' has no expression of wish" in md)
check("no-trust life policy noticed",
      "Term life policy' has no nomination / trust" in md)
check("completeness section framed as plan-readiness, no judgment "
      "calls (consolidation rule removed pending age data)",
      "## Crisis-plan completeness" in md
      and "consider consolidating" not in md)
check("recorded will/LPA/wishes/funeral produce no gaps",
      "No will recorded" not in md and "No LPA recorded" not in md
      and "wishes tab is empty" not in md
      and "No pre-arranged funeral" not in md)

print("Gaps self-clear / appear as the workbook changes")
model = read_workbook(SAMPLE)
base_gaps = derive_gaps(model)
copy = SCRATCH / "no-will.xlsx"
shutil.copyfile(SAMPLE, copy)
from openpyxl import load_workbook  # noqa: E402

wb = load_workbook(copy)
ws = wb["key_documents"]
ws.cell(row=2, column=1).value = "passport"   # the will row vanishes
wb.save(copy)
gaps2 = derive_gaps(read_workbook(copy))
check("removing the will row raises the make-a-will gap",
      any("No will recorded" in g for g in gaps2)
      and not any("No will recorded" in g for g in base_gaps))

print("Done status renders ticked")
done_copy = SCRATCH / "done.xlsx"
shutil.copyfile(SAMPLE, done_copy)
wb = load_workbook(done_copy)
ws = wb["preparation"]
header = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}
ws.cell(row=2, column=header["status"]).value = "done"
wb.save(done_copy)
md_done = render_preparation(read_workbook(done_copy))
check("done action shows [x]",
      "- [x] **NOW** Make/update the wills" in md_done)

print("Per-owner preparation lists (work shared between people)")
check("health line reports personal lists",
      "2 personal list(s)" in r.stdout)
pat = (out / "preparation-pat.md").read_text(encoding="utf-8")
alex = (out / "preparation-alex.md").read_text(encoding="utf-8")
check("each owner gets their own share",
      "Make/update the wills" in pat and "Interview candidates" in alex
      and "Interview candidates" not in pat
      and "Make/update the wills" not in alex)
check("personal list points back at the master",
      "full picture" in pat and "preparation.md" in pat)
check("unowned and not-needed items stay master-only",
      "Remove loose rugs" not in pat + alex
      and "Consider a stairlift" not in pat + alex)
ghost = out / "preparation-ghost.md"
ghost.write_text("stale", encoding="utf-8")
subprocess.run(
    [sys.executable, str(REPO / "tools" / "generate.py"), "-f",
     str(SAMPLE), "--date", "2030-01-15", "--out", str(out)],
    capture_output=True, text=True)
check("stale personal preparation lists cleaned", not ghost.exists())

print("Strictly local: never tasks, never calendar")
checklist = (out / "checklist.md").read_text(encoding="utf-8")
ics = (out / "plan.ics").read_bytes().decode("utf-8")
preview = (out / "google-preview.md").read_text(encoding="utf-8")
for term in ("Interview candidates", "stairlift", "preparation/"):
    check(f"'{term}' appears in no crisis artifact",
          term not in checklist and term not in ics
          and term not in preview)

print()
if failures:
    raise SystemExit(f"FAILED: {failures}")
print(f"All step-5.3 checks passed ({len(base_gaps)} gaps on the "
      "sample).")
