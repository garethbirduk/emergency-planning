"""Regression suite: crisis roles are an enum (roles.py).

The bug this locks down: ``crisis_roles`` was a free-text cell, so
"organizer" made the role held by nobody, the fallback quietly handed its
tasks to the executor, and the plan still read as correct. Roles now come
from one list; unrecognised values are a visible note, never a shrug.
"""

import re
import sys
import tempfile
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SCRATCH = Path(tempfile.mkdtemp(prefix="crisisplan-verifyroles-"))
sys.path.insert(0, str(REPO / "tools"))

from openpyxl import Workbook                                    # noqa: E402
import make_template                                             # noqa: E402
import roles                                                     # noqa: E402
import wizard                                                    # noqa: E402
from outputs import render_markdown                              # noqa: E402
from reader import read_workbook                                 # noqa: E402
from tasks import derive_tasks                                   # noqa: E402

failures = []


def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}"
          + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(name)


def workbook_with(crisis_roles_cell: str, path: Path) -> Path:
    """Smallest workbook the engine will read: a planholder, two helpers
    (one carrying the cell under test) and one sole bank account."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("household")
    ws.append(["name", "relationship", "role", "crisis_roles", "invited?",
               "region", "notes"])
    ws.append(["Pat Test", "self", "planholder", "", "Y", "", ""])
    ws.append(["Alex Test", "spouse", "executor", "admin", "Y", "", ""])
    ws.append(["Chris Test", "sibling", "helper", crisis_roles_cell, "Y",
               "", ""])
    bank = wb.create_sheet("banks")
    bank.append(["bank", "account_label", "joint?", "owner",
                 "location_of_details", "notes"])
    bank.append(["Test Bank plc", "Current account", "sole", "Pat",
                 "Home safe", ""])
    wb.save(path)
    return path


def plan_for(cell: str, name: str):
    """-> (tasks, notes) from a workbook whose helper holds `cell`."""
    path = workbook_with(cell, SCRATCH / f"{name}.xlsx")
    return derive_tasks(read_workbook(path), date(2030, 1, 15))


print("parse: the enum, its spellings, and what it refuses to guess")
check("plain list parses in order",
      roles.parse("admin, organiser") == (["admin", "organiser"], []))
check("case, spacing and separators tolerated",
      roles.parse(" Admin ;LOCAL\nmedical/comms ")[0]
      == ["admin", "local", "medical", "comms"])
check("US spelling accepted as the same role",
      roles.parse("Organizer")[0] == ["organiser"])
check("long forms accepted (administrator, communications)",
      roles.parse("administrator, communications")[0] == ["admin", "comms"])
check("duplicates collapse",
      roles.parse("admin, admin, ADMIN")[0] == ["admin"])
check("an unknown value is reported verbatim, not guessed",
      roles.parse("admin, hlper")== (["admin"], ["hlper"]))
check("nothing in, nothing out", roles.parse("") == ([], [])
      and roles.parse(None) == ([], []))
check("format_roles canonicalises order and drops the unknown",
      roles.format_roles(["comms", "hlper", "admin"]) == "admin, comms")
check("format_roles takes a whole cell as one item",
      roles.format_roles(["local, admin"]) == "admin, local")

print("The engine: a misspelt role is a visible note, not a silent gap")
tasks, notes = plan_for("organizer", "spelling-variant")
check("a spelling variant still holds the role (no warning)",
      not [n for n in notes if "not one of" in n]
      and any(t.assignee == "Chris Test" and t.role == "organiser"
              for t in tasks))

tasks, notes = plan_for("hlper, medical", "unknown-token")
warn = [n for n in notes if "not one of" in n]
check("an unrecognised role produces exactly one note", len(warn) == 1,
      warn[0] if warn else "no note")
check("the note names the person, the value and the valid set",
      bool(warn) and "Chris Test" in warn[0] and "'hlper'" in warn[0]
      and all(r in warn[0] for r in roles.ROLE_NAMES))
check("the roles alongside it still work",
      any(t.assignee == "Chris Test" and t.role == "medical"
          for t in tasks))
md = render_markdown(tasks, notes, "test.xlsx", date(2030, 1, 15))
check("the warning reaches the checklist, where a human will see it",
      "not one of" in md and "'hlper'" in md)

print("The wizard: an enum can't be typed wrong")
page = wizard.render_form()
check("crisis_roles rendered as tick-boxes, not a text field",
      'name="household__IDX__crisis_roles" value="admin"' in page
      and 'name="household__IDX__crisis_roles" value="comms"' in page)
check("every role offered, with the words that explain it",
      all(f'value="{r}"' in page and b in page
          for r, b in roles.CRISIS_ROLES))
check("no free-text input left for the column",
      '<input name="household__IDX__crisis_roles"' not in page)

posted = {"household__0__name": ["Chris Test"],
          "household__0__crisis_roles": ["medical", "admin"]}
rows = wizard.build_rows(posted)
check("ticked boxes join into one canonical cell",
      rows["household"][0][3] == "admin, medical",
      rows["household"][0][3])
check("ticks alone don't invent a row",
      "household" not in wizard.build_rows(
          {"household__0__crisis_roles": ["admin"]}))

print("The editor: loading a workbook ticks the boxes, hides nothing")
path = workbook_with("organizer, hlper", SCRATCH / "editable.xlsx")
err, _plan_for, _out, prefill, _maxids = wizard.load_prefill(str(path))
helper = [r for r in prefill["household"]
          if r.get("name") == "Chris Test"][0]
check("workbook loads for editing", err is None, err or "")
check("recognised value normalised so the box ticks",
      helper["crisis_roles"] == "organiser", helper["crisis_roles"])
check("unrecognised value carried through to be shown, not deleted",
      helper["_badroles"] == ["hlper"], str(helper.get("_badroles")))
check("the form knows how to show it",
      "rolewarn" in page and "not a known role" in wizard.SCRIPT_JS)

print("One source of truth")
instructions = "\n".join(str(t) for t, _bold in make_template.INSTRUCTIONS)
check("template instructions list exactly the roles from roles.py",
      all(f"{r} (" in instructions for r in roles.ROLE_NAMES))
check("the engine keeps no second copy of the list",
      not re.search(r'"(admin|organiser)"\s*,\s*"(organiser|local)"',
                    (REPO / "tools" / "tasks.py").read_text(
                        encoding="utf-8")))

print()
if failures:
    raise SystemExit(f"FAILED: {len(failures)} check(s): "
                     + "; ".join(failures))
print(f"All crisis-role checks passed ({len(roles.ROLE_NAMES)} roles).")
