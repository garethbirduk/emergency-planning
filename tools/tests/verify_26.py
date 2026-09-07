"""Regression suite from build step 2.6: friends-facing funeral artifact
(SPEC §7 — distinct audience, distinct artifact).

Sample funeral row: Example Funerals Ltd, Sampletown Crematorium,
service 2030-01-24 14:00, wake at The Sample Arms. The artifact must
carry ONLY the share fields — never internal funeral fields (package,
prepaid, contact hint, paperwork pointer) and never any other tab's
rows.
"""

import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, time
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SAMPLE = REPO / "samples" / "pat.sample.xlsx"
SCRATCH = Path(tempfile.mkdtemp(prefix="crisisplan-verify26-"))
sys.path.insert(0, str(REPO / "tools"))

failures = []


def check(name, ok, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if not ok:
        failures.append(name)


def run_cli(src, out_dir):
    r = subprocess.run(
        [sys.executable, str(REPO / "tools" / "generate.py"), "-f", str(src),
         "--date", "2030-01-15", "--out", str(out_dir)],
        capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr)
        raise SystemExit("CLI failed")
    return r.stdout


def edit_funeral(dest, **cells):
    """Copy the sample and overwrite named funeral-tab cells on row 2."""
    from openpyxl import load_workbook
    shutil.copyfile(SAMPLE, dest)
    wb = load_workbook(dest)
    ws = wb["funeral"]
    header = {c.value: i + 1 for i, c in enumerate(ws[1]) if c.value}
    for name, value in cells.items():
        ws.cell(row=2, column=header[name]).value = value
    wb.save(dest)
    return dest


print("Artifact generated from the dated sample funeral row")
out = SCRATCH / "run-a"
stdout = run_cli(SAMPLE, out)
md_path = out / "funeral-details.md"
ics_path = out / "funeral-details.ics"
check("both files written", md_path.exists() and ics_path.exists())
check("health summary reports the artifact",
      "friends artifact" in stdout and "share with the wider circle"
      in stdout)
md = md_path.read_text(encoding="utf-8")
ics = ics_path.read_bytes().decode("utf-8").replace("\r\n ", "")

print("Content: share fields only")
check("date, time, venue, wake, provider all present",
      "Thursday 24 January 2030" in md and "14:00" in md
      and "Sampletown Crematorium" in md
      and "The Sample Arms, Sampletown" in md
      and "Example Funerals Ltd" in md)
internal = ["cremation package", "prepaid", "paid for",
            "funeral folder", "home safe"]
check("no internal funeral fields leak (package/prepaid/pointers)",
      not any(t in md.lower() or t in ics.lower() for t in internal),
      ", ".join(t for t in internal
                if t in md.lower() or t in ics.lower()))
other_tabs = ["Pat", "Alex", "Chris", "Example Bank",
              "Sampletown Building Society", "Example Street",
              "Password manager", "pension", "solicitor"]
check("no other tab's rows leak into the artifact",
      not any(t in md or t in ics for t in other_tabs),
      ", ".join(t for t in other_tabs if t in md or t in ics))

print("Single-event .ics")
check("exactly one VEVENT", ics.count("BEGIN:VEVENT") == 1)
check("stable non-name UID", "UID:funeral/1/service" in ics)
check("timed start/end (14:00–15:00)",
      "DTSTART:20300124T140000" in ics
      and "DTEND:20300124T150000" in ics)
check("summary + location set",
      "SUMMARY:Funeral service" in ics
      and "LOCATION:Sampletown Crematorium" in ics)

print("Reruns idempotent")
out_b = SCRATCH / "run-b"
run_cli(SAMPLE, out_b)
check("funeral-details.md byte-identical",
      md_path.read_bytes() == (out_b / "funeral-details.md").read_bytes())
check("funeral-details.ics byte-identical",
      ics_path.read_bytes()
      == (out_b / "funeral-details.ics").read_bytes())

print("No service_date -> no artifact (the re-run flow)")
undated = edit_funeral(SCRATCH / "undated.xlsx", service_date=None)
out_c = SCRATCH / "run-c"
stdout = run_cli(undated, out_c)
check("no files written",
      not (out_c / "funeral-details.md").exists()
      and not (out_c / "funeral-details.ics").exists())
check("health says why, visibly",
      "no service_date yet" in stdout)

print("Empty funeral tab -> no artifact, base task back")
empty = SCRATCH / "empty.xlsx"
shutil.copyfile(SAMPLE, empty)
from openpyxl import load_workbook  # noqa: E402

wb = load_workbook(empty)
ws = wb["funeral"]
for i, cell in enumerate(ws[1]):
    if cell.value and cell.value not in ("id", "online_visibility",
                                         "ai_visible?"):
        ws.cell(row=2, column=i + 1).value = None
wb.save(empty)
out_d = SCRATCH / "run-d"
stdout = run_cli(empty, out_d)
check("no files and 'no funeral tab data' reported",
      not (out_d / "funeral-details.md").exists()
      and "no funeral tab data" in stdout)

print("No service_time -> all-day event")
allday = edit_funeral(SCRATCH / "allday.xlsx", service_time=None)
out_e = SCRATCH / "run-e"
run_cli(allday, out_e)
ics_e = (out_e / "funeral-details.ics").read_bytes().decode("utf-8")
check("all-day DTSTART/DTEND",
      "DTSTART;VALUE=DATE:20300124" in ics_e
      and "DTEND;VALUE=DATE:20300125" in ics_e)
check("md omits the time line",
      "**Time:**" not in (out_e / "funeral-details.md")
      .read_text(encoding="utf-8"))

print("Excel-native date/time cells work too")
native = edit_funeral(SCRATCH / "native.xlsx",
                      service_date=datetime(2030, 1, 24),
                      service_time=time(14, 0))
out_f = SCRATCH / "run-f"
run_cli(native, out_f)
check("same event from datetime/time cells",
      "DTSTART:20300124T140000"
      in (out_f / "funeral-details.ics").read_bytes().decode("utf-8"))

from outputs import _as_date, _as_time  # noqa: E402

check("junk service_date treated as unknown",
      _as_date("sometime in spring") is None
      and _as_time("mid afternoon") is None
      and _as_date(" 2030-01-24 ") == date(2030, 1, 24))

print()
if failures:
    raise SystemExit(f"FAILED: {failures}")
print("All step-2.6 checks passed.")
