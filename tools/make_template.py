#!/usr/bin/env python3
"""Template workbook generator (build step 1.1).

Builds ``samples/pat.sample.xlsx``: every module tab from SPEC §3, an
instructions tab, protected sheets with only input cells unlocked, dropdown
validation on constrained fields, and clearly-fake dummy rows.

Dummy data only — see ../PRIVACY.md. The real workbook is a private copy of
this template, filled in by hand outside the repo, never tracked, and never
shown to an AI. Pointers, never secrets: no account numbers, passwords, or
reference numbers belong in any workbook column — only where to find them.

Usage:
  python tools/make_template.py                       # rebuild the sample
  python tools/make_template.py --blank --out PATH    # blank template for
                                                      # a new household
                                                      # (step 4.1)
"""

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from roles import CRISIS_ROLES

OUT_PATH = Path(__file__).resolve().parent.parent / "samples" / "pat.sample.xlsx"

# Rows 2..INPUT_ROWS+1 are unlocked for input on every data tab.
INPUT_ROWS = 200
# Rows that get a pre-filled `id` (SPEC §13) — human-owned after that.
ID_PREFILL_ROWS = 30

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4F4F6F")
UNLOCKED = Protection(locked=False)

# Constrained-field option lists. "strict" lists block other values (Y/N-style
# facts); non-strict lists suggest via dropdown but allow free text.
YN = ["Y", "N"]
JOINT = ["sole", "joint"]
VISIBILITY = ["full", "generic", "none"]

# Each tab: list of (header, width, options, strict) column tuples + dummy rows.
# Unknown columns added later by the human are fine — the reader ignores them.
TABS = {
    "household": {
        "columns": [
            ("name", 22, None, False),
            ("relationship", 20, None, False),
            # Legal standing in the crisis, not the relationship (that has
            # its own column) and not who does the work (crisis_roles).
            # planholder = the person the plan is for; attorney_LPA holds
            # an LPA (ends at death); next_of_kin = who would register the
            # death. Blank is fine for everyone else.
            ("role", 18, ["planholder", "executor", "attorney_LPA",
                          "next_of_kin"], False),
            # Functional crisis roles (SPEC §14), comma-separated. The
            # values are an enum — roles.py owns the list; unrecognised
            # ones are warned about, never silently unheld. No dropdown:
            # the cell holds several, and Excel's list validation would
            # replace them all with the one you picked.
            ("crisis_roles", 26, None, False),
            ("invited?", 10, YN, True),
            ("region", 18, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            # `comms` is deliberately held by nobody — exercises the
            # visible-fallback path (AC-R1).
            ("Pat Sample", "self (plan holder)", "planholder", "", "Y",
             "England & Wales", "The person this plan is for"),
            ("Alex Sample", "spouse", "executor", "admin, organiser", "Y",
             "England & Wales", ""),
            ("Chris Sample", "sibling", "", "local, medical", "Y",
             "", ""),
            ("Jordan Sample", "adult child", "next_of_kin", "", "Y",
             "Scotland", "Lives away — phone first"),
        ],
    },
    "banks": {
        "columns": [
            ("bank", 24, None, False),
            ("account_label", 26, None, False),
            ("joint?", 10, JOINT, True),
            ("owner", 14, None, False),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            # Dummy rows exercise all visibility levels (trailing columns:
            # id [pre-filled later], online_visibility, ai_visible?).
            ("Example Bank plc", "Current account", "joint", "",
             "Folder 'Money' in the home safe", "", None, "full", "N"),
            ("Example Bank plc", "Savings account", "sole", "Pat",
             "Password manager entry 'Example Bank'", "", None, "", "Y"),
            ("Another Bank Ltd", "Current account", "sole", "Alex",
             "Statements drawer in the study", ""),
            ("Example Private Savings Bank", "Reserve account", "sole",
             "Pat", "Statements folder in the home safe",
             "Kept off the shared calendar entirely", None, "none", "N"),
        ],
    },
    "building_societies": {
        "columns": [
            ("society", 28, None, False),
            ("account_label", 26, None, False),
            ("joint?", 10, JOINT, True),
            ("owner", 14, None, False),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Sampletown Building Society", "Instant access saver", "sole",
             "Pat", "Passbook in the home safe", "", None, "generic", "Y"),
        ],
    },
    "cash_isas": {
        "columns": [
            ("provider", 26, None, False),
            ("account_label", 26, None, False),
            ("owner", 14, None, False),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Example Bank plc", "Cash ISA", "Pat", "Password manager", ""),
        ],
    },
    "shares_isas": {
        "columns": [
            ("provider", 26, None, False),
            ("account_label", 26, None, False),
            ("owner", 14, None, False),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Example Investments", "Stocks & shares ISA", "Alex",
             "Password manager entry 'Example Investments'", ""),
            ("Example Shares Platform", "Stocks & shares ISA", "Pat",
             "Login pointer in the password manager", "", None, "full",
             "N"),
        ],
    },
    "share_portfolios": {
        "columns": [
            ("registrar_or_platform", 26, None, False),
            ("holding_label", 30, None, False),
            ("joint?", 10, JOINT, True),
            ("owner", 14, None, False),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Example Registrars", "Utility Co shares (certificated)", "sole",
             "Pat", "Certificates in the home safe", ""),
            ("Example Platform Nominees", "Blue-chip holding", "joint", "",
             "Platform login in the password manager", ""),
        ],
    },
    "pensions": {
        "columns": [
            ("provider", 26, None, False),
            ("pension_label", 24, None, False),
            # A SIPP counts as "personal"; final_salary (defined benefit)
            # beats "workplace" — it's the one with a survivor's pension.
            ("type", 14, ["workplace", "final_salary", "personal", "state",
                          "annuity"], False),
            ("owner", 14, None, False),
            ("nomination_in_place?", 18, YN, True),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Example Pensions Ltd", "Workplace pension", "workplace", "Pat",
             "Y", "Annual statement folder in the study", "", None,
             "full", "N"),
            ("State pension", "State pension", "state", "Pat", "N", "", ""),
            ("Example Annuity Co", "Lifetime annuity", "annuity", "Pat",
             "N", "Annuity schedule in the study", ""),
        ],
    },
    "property": {
        "columns": [
            ("property_label", 20, None, False),
            ("address", 34, None, False),
            ("joint?", 10, JOINT, True),
            ("owner", 14, None, False),
            ("tenancy", 18, ["joint_tenants", "tenants_in_common", "n/a"],
             False),
            ("mortgage?", 10, YN, True),
            ("insurer", 22, None, False),
            ("location_of_deeds", 32, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Family home", "1 Example Street, Sampletown SA1 1AA", "joint",
             "", "joint_tenants", "N", "Example Insurance Co",
             "Deeds with solicitor (Example & Co)", "", None, "none", "N"),
            ("Rental flat", "2 Sample Road, Sampletown SA2 2BB", "sole",
             "Pat", "n/a", "Y", "Example Insurance Co",
             "Land Registry (registered)", ""),
            ("Holiday cottage", "3 Sample Lane, Otherville OT1 1ZZ",
             "joint", "", "tenants_in_common", "N", "Sample Cover Ltd",
             "Land Registry (registered)", "Owned with Pat's sibling",
             None, "full", "N"),
        ],
    },
    "other_assets": {
        "columns": [
            ("asset_label", 26, None, False),
            ("type", 14, ["vehicle", "valuable", "collection", "other"],
             False),
            ("joint?", 10, JOINT, True),
            ("owner", 14, None, False),
            ("location", 24, None, False),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Car — blue hatchback", "vehicle", "sole", "Pat", "Driveway",
             "V5C in the filing cabinet", "", None, "full", "N"),
            ("Wedding jewellery", "valuable", "joint", "", "Home safe",
             "Valuation kept with insurer paperwork", ""),
            ("Motorbike — red classic", "vehicle", "joint", "", "Garage",
             "V5C with the vehicle paperwork", ""),
        ],
    },
    "funeral_contacts": {
        "columns": [
            ("name", 24, None, False),
            ("relationship", 20, None, False),
            ("contact_hint", 34, None, False),
            # notify_only = must be told, but is not being invited.
            ("invite_to", 12, ["funeral", "wake", "both", "notify_only"],
             True),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Chris Sample", "sibling", "Phone in Alex's contacts", "both",
             "", None, "full", "N"),
            ("The Sample cousins", "family", "Address book in the study",
             "funeral", ""),
            ("Old colleagues group", "friends",
             "Email group — see password manager", "wake", ""),
        ],
    },
    "preparation": {
        "columns": [
            ("group", 20, None, False),
            ("item", 40, None, False),
            ("type", 14, ["action", "preference", "note"], False),
            ("owner", 16, None, False),
            ("priority", 12, ["now", "soon", "someday"], True),
            # not_needed = closed in the wizard (restorable, kept out
            # of preparation.md's active list).
            ("status", 14, ["todo", "in_progress", "done",
                            "not_needed"], True),
            ("notes", 40, None, False),
        ],
        "rows": [
            # The "before someone dies" workstream: actions, decisions
            # and living preferences. Status is tracked HERE (the tool
            # reads it, never writes it).
            ("", "Make/update the wills", "action", "Pat", "now",
             "in_progress", "Booked with Example & Co"),
            ("Arrange cleaner", "Advertise for a cleaner", "action",
             "Alex", "soon", "done", ""),
            ("Arrange cleaner", "Interview candidates", "action",
             "Alex", "soon", "todo", ""),
            ("Arrange cleaner", "Agree terms and start", "action",
             "Alex", "soon", "todo", ""),
            ("Home accessibility", "Consider a stairlift", "action",
             "", "someday", "not_needed", "Doesn't want one — see "
             "preferences"),
            ("Home accessibility", "Remove loose rugs; add grab rails",
             "action", "", "soon", "todo", ""),
            ("", "No stairlift — wants to stay mobile", "preference",
             "", "", "", "Discussed and agreed 2026"),
            ("", "Password audit — no reused or stale passwords",
             "action", "Pat", "soon", "todo",
             "Everything into the password manager"),
            ("", "Investigate residential-care options and costs",
             "action", "", "someday", "todo",
             "Just research — no decisions implied"),
        ],
    },
    "funeral": {
        "columns": [
            ("provider", 26, None, False),
            ("contact_hint", 32, None, False),
            ("package", 30, None, False),
            ("prepaid?", 10, YN, True),
            ("venue", 28, None, False),
            # Fill these in once the real details are known, then re-run:
            # they drive the separate share-with-friends artifact.
            ("wake_venue", 28, None, False),
            ("service_date", 14, None, False),
            ("service_time", 12, None, False),
            ("location_of_paperwork", 32, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            # A populated row here overrides the generic "choose a funeral
            # director" task — the funeral is pre-arranged.
            ("Example Funerals Ltd", "Number in the funeral folder",
             "Simple cremation package", "Y", "Sampletown Crematorium",
             "The Sample Arms, Sampletown", "2030-01-24", "14:00",
             "Funeral folder in the home safe", ""),
        ],
    },
    "professionals": {
        "columns": [
            ("name", 28, None, False),
            ("type", 18, ["solicitor", "accountant", "financial_adviser",
                          "other"], False),
            ("contact_hint", 34, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Example & Co Solicitors", "solicitor",
             "Office number in the address book",
             "Holds the will — see key_documents"),
            ("Sample & Partners Accountancy", "accountant",
             "Email in the address book", "Does the annual tax return"),
        ],
    },
    "life_insurance": {
        "columns": [
            ("provider", 28, None, False),
            ("policy_label", 28, None, False),
            ("type", 16, ["term", "whole_of_life", "over_50s",
                          "death_in_service", "other"], False),
            ("owner", 14, None, False),
            ("nomination_in_place?", 18, YN, True),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Example Life Assurance plc", "Term life policy", "term",
             "Pat", "N", "Policy folder in the home safe", ""),
            # Death-in-service: the classic "nobody knew it existed" asset.
            ("Sample Employer Ltd", "Death in service (4x salary)",
             "death_in_service", "Pat", "Y", "Staff handbook / HR intranet",
             ""),
        ],
    },
    "access_pointers": {
        "columns": [
            ("what", 30, None, False),
            ("where_to_find", 36, None, False),
            ("custodian", 22, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("Email account login", "Password manager (family vault)", "Alex",
             "POINTER ONLY — never the password itself"),
            ("Home safe key", "Kitchen drawer, labelled tin", "Pat", ""),
            ("Password manager emergency kit",
             "Sealed envelope with solicitor", "Example & Co Solicitors", ""),
        ],
    },
    "wishes": {
        "columns": [
            ("topic", 22, ["funeral", "burial_or_cremation", "DNACPR",
                           "advance_decision", "organ_donation", "other"],
             False),
            ("wish", 40, None, False),
            ("documented_where", 32, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("burial_or_cremation", "Cremation", "Will, clause 3", "", None,
             "full", "N"),
            ("funeral", "Simple service, family flowers only",
             "Letter of wishes kept with the will", ""),
            ("organ_donation", "Registered donor", "Donor register", ""),
            ("DNACPR", "No DNACPR in place — full treatment wanted",
             "Discussed with the GP", ""),
        ],
    },
    "key_documents": {
        "columns": [
            ("document", 20, ["will", "letter_of_wishes", "deeds",
                              "LPA_finance", "LPA_health",
                              "birth_certificate", "marriage_certificate",
                              "passport", "insurance_policy", "other"], False),
            ("location", 32, None, False),
            ("custodian", 24, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("will", "Safe at Example & Co Solicitors", "Example & Co",
             "Copy in the home safe"),
            ("deeds", "With solicitor", "Example & Co", ""),
            ("LPA_finance", "Home safe", "Alex", "Registered with OPG"),
            ("marriage_certificate", "Home safe", "Pat",
             "Needed for probate and pension claims", None, "full", "N"),
        ],
    },
    "digital": {
        "columns": [
            ("service", 26, None, False),
            ("type", 16, ["subscription", "domain", "social", "email",
                          "cloud_storage", "other"], False),
            ("action_on_death", 16, ["cancel", "memorialise", "transfer",
                                     "archive", "keep", "other"], False),
            ("owner", 14, None, False),
            ("location_of_details", 36, None, False),
            ("notes", 40, None, False),
        ],
        "rows": [
            ("example.com domain", "domain", "transfer", "Pat",
             "Registrar login in password manager", "", None, "full",
             "N"),
            ("Streaming subscription", "subscription", "cancel", "",
             "Password manager", ""),
            ("Social media profile", "social", "memorialise", "Pat", "", ""),
            ("Cloud photo storage", "cloud_storage", "archive", "Pat",
             "Password manager", ""),
            ("Example forum account", "social", "", "Pat", "",
             "No decision made yet"),
        ],
    },
}

# Crisis-chronological sheet order (SPEC §4): people, wishes, funeral,
# documents, access, contacts, professionals, then assets, then digital.
# Purely presentational — the reader is order-agnostic.
TAB_ORDER = [
    "household", "preparation",
    "wishes", "funeral", "key_documents", "access_pointers",
    "funeral_contacts", "professionals",
    "banks", "building_societies", "cash_isas", "shares_isas",
    "share_portfolios", "pensions", "property", "other_assets",
    "life_insurance",
    "digital",
]
assert set(TAB_ORDER) == set(TABS), "TAB_ORDER out of sync with TABS"

# Visibility columns (SPEC §13) appended to every data tab. `id` is the
# row's stable reference, pre-filled by build_tab; the other two control
# what external surfaces (calendar) and the optional AI lookup may see.
META_COLUMNS = [
    ("id", 6, None, False),
    ("online_visibility", 16, VISIBILITY, True),
    ("ai_visible?", 11, YN, True),
]
for _spec in TABS.values():
    _spec["columns"] = _spec["columns"] + META_COLUMNS


INSTRUCTIONS = [
    ("Crisis-plan workbook — instructions", True),
    ("", False),
    ("What this is", True),
    ("This workbook is the single source of truth for one household's crisis "
     "plan. A script reads it (and NEVER writes to it) to generate a dated "
     "task checklist and calendar when it is needed.", False),
    ("", False),
    ("How to fill it in", True),
    ("• Type only in the unlocked (white) cells below each header row. "
     "Headers are locked so the script can always find its columns.", False),
    ("• Each tab is one category. If a category doesn't apply to you, simply "
     "leave its tab empty — an empty tab means 'not applicable'.", False),
    ("• Add one row per item (per account, per property, per contact…). "
     "Fields with a dropdown expect one of the listed values.", False),
    ("• You may add extra columns or notes for your own use — the script "
     "ignores anything it doesn't recognise.", False),
    ("• 'joint?' / 'owner' matter: a sole asset follows the probate path "
     "(more steps); a joint asset follows the survivor path (fewer).", False),
    ("• 'role' (household tab) is legal standing only: planholder (who "
     "the plan is for), executor, attorney_LPA, next_of_kin (who would "
     "register the death). The relationship column says who they are to "
     "the planholder; leave role blank for everyone else.", False),
    ("• 'crisis_roles' (household tab) shares the work out — "
     + "; ".join(f"{name} ({blurb})" for name, blurb in CRISIS_ROLES)
     + ". Comma-separated, several per person is fine, and each person "
     "then gets their own checklist. These five words are the only ones "
     "the engine understands: anything else is reported as a warning in "
     "the plan and otherwise ignored, so nobody is left thinking a role "
     "is covered when it isn't. (The wizard offers them as tick-boxes.)",
     False),
    ("• funeral tab: once the real service details are known (date, time, "
     "venue, wake), fill them in and re-run — a separate share-with-friends "
     "file (funeral-details) is generated from that tab alone, safe to send "
     "to the wider circle.", False),
    ("", False),
    ("Pointers, never secrets", True),
    ("Never put account numbers, sort codes, reference numbers, passwords, "
     "PINs, or security answers in this workbook. Every 'location_of_details' "
     "or 'where_to_find' cell is a POINTER — e.g. 'password manager entry X' "
     "or 'folder in the home safe' — to where the real details live.", False),
    ("", False),
    ("Visibility columns (id / online_visibility / ai_visible?)", True),
    ("• 'id' is the row's permanent reference number (pre-filled). Never "
     "reuse a number after deleting a row — it identifies the row in "
     "calendars forever.", False),
    ("• 'online_visibility' controls what any calendar may show for this "
     "row: full = the real name, e.g. 'Notify Sampletown BS (#3)'; generic "
     "(or blank) = only e.g. 'Building society #3'; none = this row never "
     "leaves the local files at all. The local checklist always has full "
     "detail plus a legend decoding the numbers.", False),
    ("• 'ai_visible?' = Y allows the OPTIONAL lookup step to search online "
     "for this row's institution (the name only — never owner, account or "
     "locations) to add helpful links. Blank or N = never. It always shows "
     "you the exact queries and asks before going online.", False),
    ("", False),
    ("Sheet protection", True),
    ("Sheets are protected only to prevent accidents (no password). If you "
     "need to restructure, use Review → Unprotect Sheet.", False),
    ("", False),
    ("This sample file", True),
    ("Everything in this sample is fictional dummy data. Copy the file, "
     "rename it, replace the dummy rows with your own, and keep your copy "
     "OUTSIDE this repository (it must never be committed or shown to an "
     "AI assistant).", False),
]


def build_instructions(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "instructions"
    ws.column_dimensions["A"].width = 95
    for i, (text, is_heading) in enumerate(INSTRUCTIONS, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if is_heading:
            cell.font = Font(bold=True, size=14 if i == 1 else 12)
    ws.protection.sheet = True  # fully locked; nothing to input here


def build_tab(wb: Workbook, name: str, spec: dict,
              blank: bool = False) -> None:
    ws = wb.create_sheet(name)
    ws.freeze_panes = "A2"
    for col_idx, (header, width, options, strict) in enumerate(
            spec["columns"], start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = width
        head = ws.cell(row=1, column=col_idx, value=header)
        head.font = HEADER_FONT
        head.fill = HEADER_FILL
        # Unlock the input range for this column.
        for row in range(2, INPUT_ROWS + 2):
            ws.cell(row=row, column=col_idx).protection = UNLOCKED
        if options:
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(options) + '"',
                allow_blank=True,
                showErrorMessage=strict,
            )
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{INPUT_ROWS + 1}")
    for row_idx, row in enumerate([] if blank else spec["rows"], start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    # Pre-fill stable row ids (SPEC §13); the human owns them from here
    # on. Rows that already carry an id (the wizard's editor writes
    # explicit ones) keep it; empties continue from the highest used —
    # never re-issuing a number, even after deletions.
    id_col = next(i for i, (header, *_rest) in enumerate(spec["columns"],
                                                         start=1)
                  if header == "id")
    used = set()
    for row_idx in range(2, ID_PREFILL_ROWS + 2):
        value = ws.cell(row=row_idx, column=id_col).value
        if value is not None:
            try:
                used.add(int(value))
            except (TypeError, ValueError):
                pass
    next_id = max(used, default=0) + 1
    for row_idx in range(2, ID_PREFILL_ROWS + 2):
        cell = ws.cell(row=row_idx, column=id_col)
        if cell.value is None:
            cell.value = next_id
            next_id += 1
    ws.protection.sheet = True  # accident-prevention, no password


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Build a crisis-plan workbook template.")
    ap.add_argument("--out", metavar="PATH",
                    help="where to write (default: the committed sample)")
    ap.add_argument("--blank", action="store_true",
                    help="no dummy rows — a clean template for a new "
                         "household (ids pre-filled, instructions "
                         "included)")
    args = ap.parse_args()

    out_path = Path(args.out) if args.out else OUT_PATH
    repo_root = OUT_PATH.parent.parent
    resolved = out_path.resolve()
    if resolved.is_relative_to(repo_root) and not (
            resolved.is_relative_to(repo_root / "samples")
            or resolved.is_relative_to(repo_root / "private")):
        raise SystemExit(
            f"REFUSED: --out {resolved} is inside the repo working tree. "
            "A household's workbook belongs OUTSIDE the repo (see "
            "PRIVACY.md); only samples/ and private/ are allowed inside.")

    wb = Workbook()
    build_instructions(wb)
    for name in TAB_ORDER:
        build_tab(wb, name, TABS[name], blank=args.blank)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Wrote {out_path}" + (" (blank template)" if args.blank else ""))


if __name__ == "__main__":
    main()
