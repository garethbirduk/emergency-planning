#!/usr/bin/env python3
"""Tolerant workbook reader (build step 1.2).

Parses a crisis-plan workbook (the shape built by ``make_template.py``) into
a plain internal model for the task engine:

- **Presence + content** — a module counts only if its tab exists *and* has
  at least one row with a value in a recognised column (AC-C1). A present but
  empty tab means "not applicable".
- **Tolerant** — unknown tabs and unknown columns are ignored, never errors
  (AC-C2), so the human can annotate the workbook freely.
- **Read-only** — the workbook is opened in openpyxl read-only mode and is
  never written back (AC-P1). Raw cell values are used; the reader never
  depends on Excel-computed values.
- **Stable ``planKey``** — ``"<tab>/<slug>"`` derived from each row's key
  columns, so downstream outputs (checklist items, calendar events) keep the
  same identity across reruns. Duplicate slugs within a tab get ``-2``,
  ``-3``… suffixes in sheet order.

Usage:  python tools/reader.py <workbook.xlsx>   (prints a model summary)
"""

import hashlib
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

# Known module tabs: recognised columns + the columns a row's slug is built
# from. Kept in sync with make_template.py by hand; anything else in the
# workbook is ignored, not an error.
MODULE_TABS = {
    "household": {
        "columns": ["name", "relationship", "role", "crisis_roles",
                    "invited?", "region", "notes"],
        "key_columns": ["name"],
    },
    "banks": {
        "columns": ["bank", "account_label", "joint?", "owner",
                    "location_of_details", "notes"],
        "key_columns": ["bank", "account_label"],
    },
    "building_societies": {
        "columns": ["society", "account_label", "joint?", "owner",
                    "location_of_details", "notes"],
        "key_columns": ["society", "account_label"],
    },
    "cash_isas": {
        "columns": ["provider", "account_label", "owner",
                    "location_of_details", "notes"],
        "key_columns": ["provider", "account_label"],
    },
    "shares_isas": {
        "columns": ["provider", "account_label", "owner",
                    "location_of_details", "notes"],
        "key_columns": ["provider", "account_label"],
    },
    "share_portfolios": {
        "columns": ["registrar_or_platform", "holding_label", "joint?",
                    "owner", "location_of_details", "notes"],
        "key_columns": ["registrar_or_platform", "holding_label"],
    },
    "pensions": {
        "columns": ["provider", "pension_label", "type", "owner",
                    "nomination_in_place?", "location_of_details", "notes"],
        "key_columns": ["provider", "pension_label"],
    },
    "property": {
        "columns": ["property_label", "address", "joint?", "owner", "tenancy",
                    "mortgage?", "insurer", "location_of_deeds", "notes"],
        "key_columns": ["property_label"],
    },
    "other_assets": {
        "columns": ["asset_label", "type", "joint?", "owner", "location",
                    "location_of_details", "notes"],
        "key_columns": ["asset_label"],
    },
    "preparation": {
        "columns": ["group", "item", "type", "owner", "priority",
                    "status", "notes"],
        "key_columns": ["group", "item"],
    },
    "funeral_contacts": {
        "columns": ["name", "relationship", "contact_hint", "invite_to",
                    "notes"],
        "key_columns": ["name"],
    },
    "funeral": {
        "columns": ["provider", "contact_hint", "package", "prepaid?",
                    "venue", "wake_venue", "service_date", "service_time",
                    "location_of_paperwork", "notes"],
        "key_columns": ["provider"],
    },
    "professionals": {
        "columns": ["name", "type", "contact_hint", "notes"],
        "key_columns": ["name"],
    },
    "life_insurance": {
        "columns": ["provider", "policy_label", "type", "owner",
                    "nomination_in_place?", "location_of_details", "notes"],
        "key_columns": ["provider", "policy_label"],
    },
    "access_pointers": {
        "columns": ["what", "where_to_find", "custodian", "notes"],
        "key_columns": ["what"],
    },
    "wishes": {
        "columns": ["topic", "wish", "documented_where", "notes"],
        "key_columns": ["topic"],
    },
    "key_documents": {
        "columns": ["document", "location", "custodian", "notes"],
        "key_columns": ["document"],
    },
    "digital": {
        "columns": ["service", "type", "action_on_death", "owner",
                    "location_of_details", "notes"],
        "key_columns": ["service"],
    },
}

# Visibility columns (SPEC §13), present on every data tab. They are
# metadata about a row, not data: a row counts as populated only if a
# non-meta column has a value.
META_COLUMNS = ["id", "online_visibility", "ai_visible?"]
for _spec in MODULE_TABS.values():
    _spec["columns"] = _spec["columns"] + META_COLUMNS

# Known tabs that are not modules (no rows to read).
NON_MODULE_TABS = {"instructions"}

SLUG_MAX_LEN = 60


@dataclass
class Row:
    """One populated data row: recognised fields + its stable plan key."""
    tab: str
    plan_key: str
    fields: dict
    # Visibility metadata (SPEC §13); absent columns give the safe defaults.
    row_id: str | None = None      # human-owned stable reference ("3")
    visibility: str = "generic"    # full | generic | none
    ai_visible: bool = False
    label: str = ""                # human label (key columns) for legends

    @property
    def external_ref(self) -> str:
        """Non-identifying row reference: the id, or a short opaque hash of
        the slug when the id is blank (stable, but fill the id in)."""
        if self.row_id:
            return self.row_id
        slug = self.plan_key.split("/", 1)[1]
        return hashlib.sha1(slug.encode("utf-8")).hexdigest()[:6]


@dataclass
class Model:
    """Everything the task engine needs, plus health-check bookkeeping."""
    source: Path
    modules: dict = field(default_factory=dict)   # tab -> [Row]; content only
    empty_tabs: list = field(default_factory=list)    # present, no data rows
    missing_tabs: list = field(default_factory=list)  # known tab absent
    unknown_tabs: list = field(default_factory=list)  # ignored (AC-C2)


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return slug[:SLUG_MAX_LEN].rstrip("-") or "row"


def _clean(value):
    """Normalise a raw cell value: strip strings, blank -> None."""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _read_tab(ws, tab: str, spec: dict) -> list:
    """Read one module tab into Rows, ignoring unknown columns."""
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    known = set(spec["columns"])
    # Column index -> recognised column name; unknown headers dropped.
    col_map = {}
    for idx, cell in enumerate(header):
        name = _clean(cell)
        if isinstance(name, str) and name.lower() in known:
            col_map[idx] = name.lower()

    rows = []
    slug_counts = {}
    for raw in rows_iter:
        fields = {}
        for idx, name in col_map.items():
            fields[name] = _clean(raw[idx]) if idx < len(raw) else None
        if not any(v is not None for k, v in fields.items()
                   if k not in META_COLUMNS):
            continue  # blank row (a pre-filled id alone is not data)
        # Key parts joined for the slug/label — dropping blanks and
        # (case-insensitive) duplicates, so provider "Zurich" + label
        # "Zurich" keys as "zurich", never "zurich-zurich".
        key_parts, seen = [], set()
        for column in spec["key_columns"]:
            value = fields.get(column)
            part = str(value).strip() if value is not None else ""
            if part and part.lower() not in seen:
                seen.add(part.lower())
                key_parts.append(part)
        slug = slugify(" ".join(key_parts)) if key_parts else "row"
        n = slug_counts[slug] = slug_counts.get(slug, 0) + 1
        if n > 1:
            slug = f"{slug}-{n}"
        raw_id = fields.get("id")
        if isinstance(raw_id, float) and raw_id.is_integer():
            raw_id = int(raw_id)
        vis = fields.get("online_visibility")
        vis = vis.strip().lower() if isinstance(vis, str) else ""
        rows.append(Row(
            tab=tab, plan_key=f"{tab}/{slug}", fields=fields,
            row_id=str(raw_id) if raw_id is not None else None,
            # Unrecognised values redact (the safe direction).
            visibility=vis if vis in ("full", "none") else "generic",
            ai_visible=str(fields.get("ai_visible?") or "").strip().upper()
            == "Y",
            label=" — ".join(key_parts) or slug))
    return rows


def read_workbook(path) -> Model:
    """Parse the workbook read-only into a Model. Never writes (AC-P1)."""
    path = Path(path)
    model = Model(source=path)
    wb = load_workbook(path, read_only=True)
    try:
        present = set(wb.sheetnames)
        for tab in wb.sheetnames:
            if tab not in MODULE_TABS and tab not in NON_MODULE_TABS:
                model.unknown_tabs.append(tab)
        for tab, spec in MODULE_TABS.items():
            if tab not in present:
                model.missing_tabs.append(tab)
                continue
            rows = _read_tab(wb[tab], tab, spec)
            if rows:
                model.modules[tab] = rows  # exists AND has data (AC-C1)
            else:
                model.empty_tabs.append(tab)
    finally:
        wb.close()
    return model


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("usage: python tools/reader.py <workbook.xlsx>")
    model = read_workbook(sys.argv[1])
    print(f"Source: {model.source}")
    for tab, rows in model.modules.items():
        print(f"\n{tab} ({len(rows)} row{'s' if len(rows) != 1 else ''}):")
        for row in rows:
            print(f"  {row.plan_key}")
    if model.empty_tabs:
        print(f"\nEmpty tabs (not applicable): {', '.join(model.empty_tabs)}")
    if model.missing_tabs:
        print(f"Missing tabs: {', '.join(model.missing_tabs)}")
    if model.unknown_tabs:
        print(f"Unknown tabs (ignored): {', '.join(model.unknown_tabs)}")


if __name__ == "__main__":
    main()
